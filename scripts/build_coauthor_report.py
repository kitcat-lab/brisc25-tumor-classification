"""Build the coauthor discussion document as Word and PDF from its Markdown source."""
import argparse
import html
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]


def blocks(text):
    lines = text.splitlines(); i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip(): i += 1; continue
        if line.startswith('#'):
            match = re.match(r'^(#+)\s+(.*)', line)
            yield 'heading', (len(match[1]), match[2]); i += 1
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', x) for x in cells): rows.append(cells)
                i += 1
            yield 'table', rows
        elif line.startswith('```'):
            i += 1; code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            i += 1; yield 'code', '\n'.join(code)
        elif re.match(r'^\d+\. ', line):
            yield 'paragraph', line; i += 1
        elif line.startswith('- '):
            yield 'paragraph', '• ' + line[2:]; i += 1
        else:
            paragraph = [line]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|\d+\. |- )', lines[i]):
                paragraph.append(lines[i]); i += 1
            yield 'paragraph', ' '.join(paragraph)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=ROOT/'docs/BRISC25_coautores_2026-10-06.md')
    parser.add_argument('--output-base', type=Path, default=ROOT/'docs/BRISC25_coautores_2026-10-06')
    parser.add_argument('--font-dir', type=Path)
    args = parser.parse_args()
    parsed = list(blocks(args.source.read_text(encoding='utf-8')))
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    normal = document.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        document.styles[name].font.color.rgb = RGBColor.from_string('5F3F5B')
    document.core_properties.author = 'Catarina Bota'
    document.core_properties.title = 'BRISC 2025 — pontos para discutir com os coautores'
    document.core_properties.subject = 'Curadoria pHash e revisão do manuscrito'
    section.footer.paragraphs[0].text = 'BRISC 2025 · Documento de discussão · 6 de outubro de 2026'
    font, bold = 'Helvetica', 'Helvetica-Bold'
    if args.font_dir:
        pdfmetrics.registerFont(TTFont('ReportText', str(args.font_dir/'arial.ttf')))
        pdfmetrics.registerFont(TTFont('ReportBold', str(args.font_dir/'arialbd.ttf')))
        pdfmetrics.registerFontFamily('ReportText', normal='ReportText', bold='ReportBold')
        font, bold = 'ReportText', 'ReportBold'
    styles = getSampleStyleSheet()
    body = ParagraphStyle('ReportBody', fontName=font, fontSize=9.4, leading=13,
                          textColor=colors.HexColor('#222222'), spaceAfter=7)
    heading = ParagraphStyle('ReportHeading', parent=body, fontName=bold, fontSize=13,
                             leading=17, spaceBefore=14, spaceAfter=8, keepWithNext=True,
                             textColor=colors.HexColor('#5F3F5B'))
    small_heading = ParagraphStyle('ReportSubheading', parent=heading, fontSize=10.4, leading=14)
    title = ParagraphStyle('ReportTitle', parent=heading, fontSize=20, leading=25, spaceBefore=0)
    cell = ParagraphStyle('ReportCell', parent=body, fontSize=8.2, leading=11, spaceAfter=0)
    cell_header = ParagraphStyle('ReportCellHeader', parent=cell, fontName=bold)
    code_style = ParagraphStyle('ReportCode', parent=body, fontSize=8.2, leading=11,
                                backColor=colors.HexColor('#F3F1EE'), borderPadding=7)
    flow = []
    for kind, value in parsed:
        if kind == 'heading':
            level, text = value
            document.add_heading(text, level=0 if level == 1 else min(level-1, 2))
            flow.append(Paragraph(html.escape(text), title if level == 1 else heading if level == 2 else small_heading))
        elif kind == 'table':
            table = document.add_table(rows=1, cols=len(value[0])); table.style = 'Light Shading Accent 1'
            for j, text in enumerate(value[0]): table.rows[0].cells[j].text = text
            for row in value[1:]:
                cells = table.add_row().cells
                for j, text in enumerate(row): cells[j].text = text
            for row in table.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        for run in p.runs: run.font.size = Pt(9)
            document.add_paragraph()
            n = len(value[0]); width = A4[0]-100
            widths = ([width*.31, width*.69] if n == 2 else [width*.4]+[width*.6/(n-1)]*(n-1))
            data = [[Paragraph(html.escape(text), cell_header if i == 0 else cell) for text in row]
                    for i, row in enumerate(value)]
            pdf_table = Table(data, colWidths=widths, repeatRows=1, hAlign='LEFT')
            pdf_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EAE3E8')),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('LEFTPADDING', (0,0), (-1,-1), 6), ('RIGHTPADDING', (0,0), (-1,-1), 6),
                ('TOPPADDING', (0,0), (-1,-1), 6), ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ('LINEBELOW', (0,0), (-1,0), .5, colors.HexColor('#AFA2AB')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F7F5F2')])]))
            flow.extend([pdf_table, Spacer(1, 10)])
        else:
            clean = value.replace('`', '')
            document.add_paragraph(clean)
            flow.append(Paragraph(html.escape(clean).replace('\n', '<br/>'), code_style if kind == 'code' else body))
    args.output_base.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_base.with_suffix('.docx'))
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont(font, 8); canvas.setFillColor(colors.HexColor('#6B6268'))
        canvas.drawString(50, 25, 'BRISC 2025 · Documento de discussão · 6 de outubro de 2026')
        canvas.drawRightString(A4[0]-50, 25, str(doc.page)); canvas.restoreState()
    pdf = SimpleDocTemplate(str(args.output_base.with_suffix('.pdf')), pagesize=A4,
        leftMargin=50, rightMargin=50, topMargin=45, bottomMargin=45,
        author='Catarina Bota', title='BRISC 2025 — pontos para discutir com os coautores')
    pdf.build(flow, onFirstPage=footer, onLaterPages=footer)
    print('Created:', args.output_base.with_suffix('.docx'))
    print('Created:', args.output_base.with_suffix('.pdf'))


if __name__ == '__main__':
    main()
